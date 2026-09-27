"""导出：Excel / CSV / 自包含 HTML 报告（HTML 报告内联图表，不依赖外部资源）。"""
from __future__ import annotations

import csv
import html as _html
import re
import time
from pathlib import Path

from ..core.config import DATA_DIR

EXPORT_DIR = DATA_DIR / "exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

# Excel 禁止的控制字符（0x00-0x08 / 0x0B / 0x0C / 0x0E-0x1F）。
# 抓来的评论里偶尔会混进 NUL 之类的不可见字符，一旦写进工作表 openpyxl 会直接抛
# IllegalCharacterError，导致**整个 xlsx 导出失败**（只因为它落在某一条评论里）。
# 这些字符对用户没有任何意义，导出时统一剔除。
_BAD_CHAR_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _cell(v):
    """清洗单元格值：字符串里的 Excel 非法控制字符一律去掉。"""
    if isinstance(v, str):
        return _BAD_CHAR_RE.sub("", v)
    return v


def _safe(name: str) -> str:
    name = _BAD_CHAR_RE.sub("", name or "report")
    name = "".join(c for c in name if c not in '\\/:*?"<>|').strip()
    return name[:60] or "report"

try:
    from openpyxl import Workbook  # type: ignore
    from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore
    HAS_OPENPYXL = True
except Exception:  # pragma: no cover
    HAS_OPENPYXL = False


def _ts(t):
    if not t:
        return ""
    try:
        return time.strftime("%Y-%m-%d %H:%M", time.localtime(float(t)))
    except Exception:
        return ""


def _memory_ctx():
    """载入两类情感记忆：语句级记忆表 + 词条级覆盖项。"""
    from ..analysis import comment_memory, sentiment_memory
    return (comment_memory.load_memory(),
            sentiment_memory.overrides_to_lexicon(sentiment_memory.load_memory()))


def _row_sentiment(content: str, cmem: dict, overrides: dict) -> tuple[str, float, str]:
    """单条评论的情感：人工标过的语句优先，否则按词典（含词条记忆）自动判定。

    返回 (标签, 得分, 来源)，来源为「人工记忆」或「自动」。
    """
    from ..analysis import comment_memory
    from ..analysis.sentiment import score_text

    label, score, _ = score_text(content, overrides)
    ml = comment_memory.label_of(cmem, comment_memory.norm_key(content))
    if ml:
        return ml, score, "人工记忆"
    return label, score, "自动"


def to_csv(task_name: str, analysis: dict, comments: list[dict]) -> Path:
    """导出评论明细；未装 openpyxl 时 Excel 导出也走这里。"""
    from ..analysis.sentiment import classify_topics, score_text

    cmem, overrides = _memory_ctx()
    p = EXPORT_DIR / f"{_safe(task_name)}_{time.strftime('%m%d%H%M')}.csv"
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["用户", "内容", "点赞", "回复", "时间", "情感", "得分", "情感来源", "主题"])
        for c in comments:
            content = c.get("content", "")
            label, score, src = _row_sentiment(content, cmem, overrides)
            w.writerow([c.get("user_name", ""), content, c.get("like_count", 0),
                        c.get("reply_count", 0), _ts(c.get("published_at")), label, score, src,
                        "、".join(classify_topics(content))])
    return p


def to_excel(task_name: str, analysis: dict, comments: list[dict]) -> Path:
    """生成多页 Excel；未装 openpyxl 时自动降级为 CSV。"""
    if not HAS_OPENPYXL:
        return to_csv(task_name, analysis, comments)
    from ..analysis.sentiment import classify_topics, score_text

    p = EXPORT_DIR / f"{_safe(task_name)}_{time.strftime('%m%d%H%M')}.xlsx"
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="3B6FD4")

    def sheet(title, headers, rows, widths=None):
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append([_cell(v) for v in r])
        for i, wd in enumerate(widths or [18] * len(headers), start=1):
            ws.column_dimensions[get_column_letter(i)].width = wd
        ws.freeze_panes = "A2"
        return ws

    st = analysis.get("stats", {})
    ws = wb.active
    ws.title = "概览"
    ws.append(["社区舆情分析概览"])
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ["生成时间", time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(analysis.get("generated_at", time.time())))],
        ["有效评论数", st.get("total", 0)],
        ["任务总评论数", st.get("task_total", st.get("total", 0))],
        ["原始条数", st.get("raw_total", 0)],
        ["合并重复", st.get("duplicates_merged", 0)],
        ["平均长度", st.get("avg_length", 0)],
        ["正面", st.get("positive", 0)],
        ["中立", st.get("neutral", 0)],
        ["负面", st.get("negative", 0)],
        ["好评率", f"{st.get('positive_rate', 0) * 100:.1f}%"],
        ["差评率", f"{st.get('negative_rate', 0) * 100:.1f}%"],
    ]
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 30

    _lbl = {"positive": "正面", "neutral": "中立", "negative": "负面"}
    sheet("高频词条", ["词条", "权重", "出现评论数", "记忆情感标签"],
          [[k["word"], k["weight"], k["docs"], _lbl.get(k.get("memory_label"), "—")]
           for k in analysis.get("keywords", [])],
          [22, 12, 14])
    sheet("新词热梗", ["词", "出现次数", "凝固度"],
          [[n["word"], n["count"], n["cohesion"]] for n in analysis.get("new_words", [])],
          [18, 12, 12])
    sheet("常见短语", ["短语", "次数"],
          [[p_["phrase"], p_["count"]] for p_ in analysis.get("phrases", [])], [24, 10])
    sheet("吐槽点归类", ["吐槽点", "条数", "负面率", "负面", "中立", "正面", "关键词"],
          [[t["topic"], t["count"], f"{t.get('neg_rate', 0) * 100:.0f}%", t["negative"],
            t["neutral"], t["positive"], "、".join(t.get("keywords", []))]
           for t in analysis.get("topics", [])], [16, 8, 10, 8, 8, 8, 34])
    sheet("相似句聚类", ["相似条数", "代表句", "关键词", "负面", "中立", "正面", "总点赞"],
          [[c["size"], c["representative"], "、".join(c.get("keywords", [])),
            c["sentiment"]["negative"], c["sentiment"]["neutral"], c["sentiment"]["positive"],
            c["total_likes"]] for c in analysis.get("clusters", [])], [10, 60, 26, 8, 8, 8, 10])
    sheet("重点原句", ["点赞", "入选理由", "情感", "内容", "时间"],
          [[h["like_count"], h["reason"], h["sentiment"], h["content"], _ts(h.get("time"))]
           for h in analysis.get("highlights", [])], [8, 22, 10, 80, 18])

    cmem, overrides = _memory_ctx()
    detail = []
    for c in comments:
        content = c.get("content", "")
        label, score, src = _row_sentiment(content, cmem, overrides)
        detail.append([c.get("user_name", ""), content, c.get("like_count", 0),
                       c.get("reply_count", 0), _ts(c.get("published_at")), label, score, src,
                       "、".join(classify_topics(content))])
    sheet("评论明细", ["用户", "内容", "点赞", "回复", "时间", "情感", "得分", "情感来源", "主题"],
          detail, [16, 90, 8, 8, 18, 10, 8, 12, 24])

    wb.save(p)
    return p


def to_html_report(task_name: str, analysis: dict, comments: list[dict]) -> Path:
    """自包含 HTML 报告：内联 CSS 图表，双击即可打开，方便直接发给研发/老板。"""
    p = EXPORT_DIR / f"{_safe(task_name)}_{time.strftime('%m%d%H%M')}.html"
    st = analysis.get("stats", {})
    e = _html.escape

    def bar_row(label, value, total, color="#3B6FD4"):
        pct = (value / total * 100) if total else 0
        return (f'<div class="bar"><span class="bl">{e(str(label))}</span>'
                f'<span class="bt"><i style="width:{pct:.1f}%;background:{color}"></i></span>'
                f'<span class="bv">{value}</span></div>')

    kws = "".join(bar_row(k["word"], k["docs"], st.get("total", 1)) for k in analysis.get("keywords", [])[:30])
    topics = "".join(
        bar_row(t["topic"], t["count"], max(len(comments), 1),
                "#d9534f" if t.get("neg_rate", 0) > 0.5 else "#e0a33e" if t.get("neg_rate", 0) > 0.25 else "#4c9a5b")
        for t in analysis.get("topics", [])[:12])
    clusters = "".join(
        f'<div class="card"><div class="ch"><b>相似 {c["size"]} 条</b>'
        f'<span class="tag {c["rep_meta"]["sentiment"]}">{_label_cn(c["rep_meta"]["sentiment"])}</span>'
        f'<span class="mut">关键词：{e("、".join(c.get("keywords", [])))}</span></div>'
        f'<div class="rep">{e(c["representative"])}</div>'
        + "".join(f'<div class="sub">{e(s["content"])} <span class="mut">👍{s["like_count"]}</span></div>'
                  for s in c.get("samples", [])[1:4])
        + "</div>" for c in analysis.get("clusters", [])[:12])
    highs = "".join(
        f'<div class="card"><div class="ch"><span class="tag {h["sentiment"]}">{_label_cn(h["sentiment"])}</span>'
        f'<span class="reason">{e(h["reason"])}</span><span class="mut">👍{h["like_count"]}</span></div>'
        f'<div class="rep">{e(h["content"])}</div></div>' for h in analysis.get("highlights", [])[:20])

    total = max(st.get("total", 1), 1)
    pos = st.get("positive", 0)
    neu = st.get("neutral", 0)
    neg = st.get("negative", 0)
    trend = analysis.get("trend") or {}
    trend_html = ""
    if trend.get("available"):
        pts = trend["points"]
        mx = max((b["count"] for b in pts), default=1)
        bars = "".join(
            f'<div class="tcol" title="{b["date"]} 共{b["count"]}条 负面{b["negative"]}">'
            f'<div class="tbar" style="height:{max(b["count"] / mx * 100, 3):.0f}%">'
            f'<i class="neg" style="height:{b["negative"] / max(b["count"], 1) * 100:.0f}%"></i></div>'
            f'<span>{b["date"][5:]}</span></div>' for b in pts)
        trend_html = f'<div class="trend">{bars}</div>'

    task_total = int(st.get("task_total") or 0)
    # 本次实际读取条数 < 任务库中总数时，导出件必须写明口径，否则会被当成全量
    # （用 raw_total 而非去重后的 total 比较，避免误判）
    scope_note = (f"（该任务共 {task_total} 条，本次分析仅覆盖 {st.get('raw_total', 0)} 条）"
                  if task_total > int(st.get("raw_total") or 0) else "")

    doc = f"""<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">
<title>{e(task_name)} - 社区舆情报告</title><style>
*{{box-sizing:border-box}}body{{margin:0;padding:32px;background:#f5f6f8;color:#1f2328;
font-family:"Microsoft YaHei","PingFang SC",-apple-system,Segoe UI,sans-serif;line-height:1.6}}
.wrap{{max-width:1080px;margin:0 auto}}h1{{font-size:24px;margin:0 0 4px}}
.mut{{color:#7a828a;font-size:12px}}h2{{font-size:17px;margin:28px 0 12px;padding-left:10px;border-left:4px solid #3B6FD4}}
.kpis{{display:grid;grid-template-columns:repeat(5,1fr);gap:12px;margin:18px 0}}
.kpi{{background:#fff;border:1px solid #e6e8eb;border-radius:10px;padding:14px;text-align:center}}
.kpi b{{display:block;font-size:22px;color:#3B6FD4}}.kpi span{{font-size:12px;color:#7a828a}}
.pie{{display:flex;height:26px;border-radius:6px;overflow:hidden;margin:10px 0}}
.pie i{{display:block}}.legend{{font-size:12px;color:#5b636b;margin-bottom:14px}}
.bar{{display:flex;align-items:center;gap:10px;margin:6px 0;font-size:13px}}
.bl{{width:130px;text-align:right;flex:none;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}}
.bt{{flex:1;background:#eef1f5;height:16px;border-radius:4px;overflow:hidden}}
.bt i{{display:block;height:100%}}.bv{{width:46px;font-size:12px;color:#5b636b}}
.card{{background:#fff;border:1px solid #e6e8eb;border-radius:10px;padding:14px;margin:10px 0}}
.ch{{display:flex;gap:10px;align-items:center;font-size:12px;margin-bottom:8px;flex-wrap:wrap}}
.rep{{font-size:14px;color:#1f2328}}.sub{{font-size:12px;color:#6b737b;margin-top:6px;padding-left:10px;border-left:2px solid #e6e8eb}}
.tag{{padding:2px 8px;border-radius:10px;font-size:11px;color:#fff}}
.tag.positive{{background:#4c9a5b}}.tag.neutral{{background:#9aa0a6}}.tag.negative{{background:#d9534f}}
.reason{{background:#eef3fd;color:#3B6FD4;padding:2px 8px;border-radius:4px}}
.trend{{display:flex;align-items:flex-end;gap:3px;height:150px;background:#fff;border:1px solid #e6e8eb;border-radius:10px;padding:12px}}
.tcol{{flex:1;display:flex;flex-direction:column;align-items:center;justify-content:flex-end;height:100%}}
.tbar{{width:100%;background:#c9dcf7;position:relative;border-radius:3px 3px 0 0}}
.tbar .neg{{position:absolute;bottom:0;width:100%;background:#d9534f;border-radius:0 0 3px 3px}}
.tcol span{{font-size:9px;color:#8b9299;margin-top:3px;transform:rotate(-45deg)}}
.foot{{margin-top:30px;font-size:12px;color:#9aa0a6;text-align:center}}
</style></head><body><div class="wrap">
<h1>{e(task_name)} · 社区舆情报告</h1>
<div class="mut">生成时间 {time.strftime('%Y-%m-%d %H:%M:%S')} ｜ 有效样本 {st.get('total', 0)} 条{scope_note}</div>
<div class="kpis">
<div class="kpi"><b>{st.get('total', 0)}</b><span>有效评论</span></div>
<div class="kpi"><b>{st.get('positive_rate', 0) * 100:.0f}%</b><span>好评率</span></div>
<div class="kpi"><b>{st.get('negative_rate', 0) * 100:.0f}%</b><span>差评率</span></div>
<div class="kpi"><b>{len(analysis.get('clusters', []))}</b><span>相似句族群</span></div>
<div class="kpi"><b>{len(analysis.get('topics', []))}</b><span>吐槽点类别</span></div>
</div>
<div class="pie"><i style="width:{pos / total * 100:.1f}%;background:#4c9a5b"></i>
<i style="width:{neu / total * 100:.1f}%;background:#9aa0a6"></i>
<i style="width:{neg / total * 100:.1f}%;background:#d9534f"></i></div>
<div class="legend">正面 {pos}（{pos / total * 100:.1f}%）｜中立 {neu}（{neu / total * 100:.1f}%）｜负面 {neg}（{neg / total * 100:.1f}%）</div>
<h2>高频词条 Top30</h2>{kws or '<div class="mut">无</div>'}
<h2>吐槽点归类</h2>{topics or '<div class="mut">无</div>'}
<h2>声量趋势</h2>{trend_html or '<div class="mut">评论缺少时间戳，无法绘制趋势</div>'}
<h2>高频相似诉求</h2>{clusters or '<div class="mut">无</div>'}
<h2>重点舆情原句</h2>{highs or '<div class="mut">无</div>'}
<div class="foot">CommunityPulse 社区舆情雷达 · 数据仅供内部运营参考</div>
</div></body></html>"""
    p.write_text(doc, encoding="utf-8")
    return p


def _label_cn(label: str) -> str:
    return {"positive": "正面", "neutral": "中立", "negative": "负面"}.get(label, label)
